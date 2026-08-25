from decimal import Decimal
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.domain.analytics.repository import PeriodTotals, ProductSales
from app.domain.common.money import Money
from app.domain.debts.entities import Debt
from app.domain.sales.entities import Sale
from app.domain.sales.payment import PaymentMethod

_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=13)

_PAYMENT_METHOD_LABELS = {
    PaymentMethod.CASH: "Наличные",
    PaymentMethod.TRANSFER: "Перевод",
    PaymentMethod.CREDIT: "В долг",
}


def _rub(value: Decimal) -> str:
    """SQL-агрегаты (SUM) отдают Decimal с произвольным числом знаков — округляем до копеек."""
    return f"{Money.of(value).amount} ₽"


def _qty(value: Decimal, fractional: bool) -> str:
    """Целые единицы (шт, упаковка…) — без дробной части; дробные (кг, л…) — без хвостовых нулей."""
    if not fractional:
        return str(int(value))
    normalized = value.normalize()
    return format(normalized, "f") if "E" in str(normalized) else str(normalized)


_MAX_TRAILING_COLUMNS = 30  # с запасом — скрыть все колонки за пределами данных, чтобы вьюеры (Telegram) не рисовали пустую сетку


def _autosize(sheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for i in range(len(widths) + 1, len(widths) + 1 + _MAX_TRAILING_COLUMNS):
        sheet.column_dimensions[get_column_letter(i)].hidden = True


def _write_header(sheet, row: int, columns: list[str]) -> None:
    for col, title in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=col, value=title)
        cell.font = _HEADER_FONT


def render_summary_xlsx(
    shop_name: str,
    period_label: str,
    totals: PeriodTotals,
    debts_issued: Decimal,
    top_products: list[ProductSales],
    debts: list[Debt],
    counterparty_name: dict[tuple[str, int], str],
    by_warehouse: list[tuple[str, PeriodTotals]] | None = None,
    sales: list[Sale] | None = None,
    product_name: dict[int, str] | None = None,
    product_fractional: dict[int, bool] | None = None,
    seller_name: dict[int, str] | None = None,
    warehouse_name: dict[int, str] | None = None,
    client_name: dict[int, str] | None = None,
) -> bytes:
    workbook = openpyxl.Workbook()

    overview = workbook.active
    overview.title = "Итоги"
    overview["A1"] = shop_name
    overview["A1"].font = _TITLE_FONT
    overview["A2"] = period_label
    overview["A2"].font = Font(italic=True, color="666666")
    _write_header(overview, 4, ["Показатель", "Значение"])
    rows = [
        ("Продаж", totals.sales_count),
        ("Выручка", _rub(totals.revenue)),
        ("Прибыль", _rub(totals.profit)),
        ("Выдано в долг", _rub(debts_issued)),
    ]
    for i, (label, value) in enumerate(rows, start=5):
        overview.cell(row=i, column=1, value=label)
        overview.cell(row=i, column=2, value=value).alignment = Alignment(horizontal="right")
    _autosize(overview, [22, 20])

    if by_warehouse:
        by_warehouse_sheet = workbook.create_sheet("По складам")
        _write_header(by_warehouse_sheet, 1, ["Склад", "Продаж", "Выручка", "Прибыль"])
        row = 2
        for wh_name, w_totals in by_warehouse:
            by_warehouse_sheet.cell(row=row, column=1, value=wh_name)
            by_warehouse_sheet.cell(row=row, column=2, value=w_totals.sales_count).alignment = Alignment(horizontal="right")
            by_warehouse_sheet.cell(row=row, column=3, value=_rub(w_totals.revenue)).alignment = Alignment(horizontal="right")
            by_warehouse_sheet.cell(row=row, column=4, value=_rub(w_totals.profit)).alignment = Alignment(horizontal="right")
            row += 1
        by_warehouse_sheet.cell(row=row, column=1, value="Итого").font = _HEADER_FONT
        by_warehouse_sheet.cell(row=row, column=2, value=sum(t.sales_count for _, t in by_warehouse)).font = _HEADER_FONT
        by_warehouse_sheet.cell(row=row, column=3, value=_rub(sum((t.revenue for _, t in by_warehouse), Decimal(0)))).font = _HEADER_FONT
        by_warehouse_sheet.cell(row=row, column=4, value=_rub(sum((t.profit for _, t in by_warehouse), Decimal(0)))).font = _HEADER_FONT
        _autosize(by_warehouse_sheet, [26, 12, 16, 16])

    product_fractional = product_fractional or {}
    products = workbook.create_sheet("Товары")
    _write_header(products, 1, ["Товар", "Продано (кол-во)", "Выручка"])
    row = 2
    for item in top_products:
        fractional = product_fractional.get(item.product_id, True)
        products.cell(row=row, column=1, value=item.product_name)
        products.cell(row=row, column=2, value=_qty(item.quantity_sold, fractional)).alignment = Alignment(horizontal="right")
        products.cell(row=row, column=3, value=_rub(item.revenue)).alignment = Alignment(horizontal="right")
        row += 1
    products.cell(row=row, column=1, value="Итого").font = _HEADER_FONT
    products.cell(row=row, column=3, value=_rub(sum((i.revenue for i in top_products), Decimal(0)))).font = _HEADER_FONT
    _autosize(products, [32, 18, 16])

    if sales is not None:
        product_name = product_name or {}
        seller_name = seller_name or {}
        warehouse_name = warehouse_name or {}
        client_name = client_name or {}

        sales_sheet = workbook.create_sheet("Продажи")
        _write_header(
            sales_sheet, 1,
            ["Дата/время", "Чек №", "Склад", "Продавец", "Клиент", "Товар", "Кол-во", "Цена", "Сумма строки", "Оплата", "Итог чека"],
        )
        row = 2
        checks_total_sum = Decimal(0)
        seen_sale_ids: set[int] = set()
        for sale in sales:
            client = client_name.get(sale.client_id, "—") if sale.client_id else "—"
            payment_label = _PAYMENT_METHOD_LABELS.get(sale.payment_method, str(sale.payment_method))
            for line in sale.lines:
                sales_sheet.cell(row=row, column=1, value=sale.created_at.strftime("%d.%m.%Y %H:%M"))
                sales_sheet.cell(row=row, column=2, value=sale.id)
                sales_sheet.cell(row=row, column=3, value=warehouse_name.get(sale.warehouse_id, "—"))
                sales_sheet.cell(row=row, column=4, value=seller_name.get(sale.seller_id, "—"))
                sales_sheet.cell(row=row, column=5, value=client)
                sales_sheet.cell(row=row, column=6, value=product_name.get(line.product_id, "—"))
                sales_sheet.cell(row=row, column=7, value=_qty(line.quantity.value, line.quantity.unit.fractional)).alignment = Alignment(horizontal="right")
                sales_sheet.cell(row=row, column=8, value=f"{line.price.amount} ₽").alignment = Alignment(horizontal="right")
                sales_sheet.cell(row=row, column=9, value=f"{line.total.amount} ₽").alignment = Alignment(horizontal="right")
                sales_sheet.cell(row=row, column=10, value=payment_label)
                sales_sheet.cell(row=row, column=11, value=f"{sale.total.amount} ₽").alignment = Alignment(horizontal="right")
                row += 1
            if sale.id not in seen_sale_ids:
                checks_total_sum += sale.total.amount
                seen_sale_ids.add(sale.id)
        sales_sheet.cell(row=row, column=10, value="Итого").font = _HEADER_FONT
        sales_sheet.cell(row=row, column=11, value=f"{checks_total_sum} ₽").font = _HEADER_FONT
        _autosize(sales_sheet, [16, 8, 16, 16, 18, 28, 10, 12, 14, 12, 12])

    debts_sheet = workbook.create_sheet("Долги")
    _write_header(debts_sheet, 1, ["Контрагент", "Сумма", "Остаток", "Дата", "Срок"])
    row = 2
    for debt in debts:
        key = ("client", debt.client_id) if debt.client_id else ("supplier", debt.supplier_id)
        debts_sheet.cell(row=row, column=1, value=counterparty_name.get(key, "—"))
        debts_sheet.cell(row=row, column=2, value=f"{debt.amount.amount} ₽").alignment = Alignment(horizontal="right")
        debts_sheet.cell(row=row, column=3, value=f"{debt.balance.amount} ₽").alignment = Alignment(horizontal="right")
        debts_sheet.cell(row=row, column=4, value=debt.issued_at.strftime("%d.%m.%Y"))
        debts_sheet.cell(row=row, column=5, value=debt.due_date.strftime("%d.%m.%Y") if debt.due_date else "—")
        row += 1
    debts_sheet.cell(row=row, column=1, value="Итого").font = _HEADER_FONT
    debts_sheet.cell(row=row, column=2, value=f"{sum((d.amount.amount for d in debts), Decimal(0))} ₽").font = _HEADER_FONT
    debts_sheet.cell(row=row, column=3, value=f"{sum((d.balance.amount for d in debts), Decimal(0))} ₽").font = _HEADER_FONT
    _autosize(debts_sheet, [28, 14, 14, 12, 12])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
